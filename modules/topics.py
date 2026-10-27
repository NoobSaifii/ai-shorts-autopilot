import logging
import socket
import ssl
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import openpyxl
import pandas as pd

logger = logging.getLogger("pipeline.topics")

_write_lock = threading.Lock()


@dataclass
class TopicRow:
    number: int
    topic: str
    series_name: str
    episode_no: int
    priority: str
    status: str


def _safe_int(val, default=0) -> int:
    s = str(val if val is not None else "")
    if s in ("nan", "", "None"):
        return default
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return default


def _col_letter(idx: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA (for A1 notation)."""
    s = ""
    idx += 1
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


# ── Excel backend (local .xlsx — the original behaviour) ───────────────────────
class _ExcelBackend:
    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Topics file not found: {file_path}")

    def get_sheet_names(self) -> List[str]:
        wb = openpyxl.load_workbook(self.file_path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    def get_pending_topics(self, sheet_name: str = "Sheet1") -> List[TopicRow]:
        df = pd.read_excel(self.file_path, sheet_name=sheet_name, dtype=str)
        df.columns = [str(c).strip().lower() for c in df.columns]

        if "status" not in df.columns:
            raise ValueError(f"Sheet '{sheet_name}' must have a 'status' column")

        pending = df[df["status"].str.strip().str.lower() == "pending"].copy()

        if "#" in pending.columns:
            pending = pending.sort_values("#", key=lambda x: pd.to_numeric(x, errors="coerce"))

        rows = []
        for _, row in pending.iterrows():
            topic = str(row.get("topic", "")).strip()
            if not topic:
                continue
            rows.append(
                TopicRow(
                    number=_safe_int(row.get("#", 0)),
                    topic=topic,
                    series_name=str(row.get("series_name", "")).strip(),
                    episode_no=_safe_int(row.get("episode_no", 0)),
                    priority=str(row.get("priority", "")).strip(),
                    status=str(row.get("status", "")).strip(),
                )
            )
        return rows

    def get_next_pending(self, sheet_name: str = "Sheet1") -> Optional[TopicRow]:
        rows = self.get_pending_topics(sheet_name=sheet_name)
        return rows[0] if rows else None

    def pending_count(self, sheet_name: str = "Sheet1") -> int:
        return len(self.get_pending_topics(sheet_name=sheet_name))

    def mark_done(self, topic: str, sheet_name: str = "Sheet1"):
        with _write_lock:
            wb = openpyxl.load_workbook(self.file_path)
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet '{sheet_name}' not found in Excel file")
            ws = wb[sheet_name]
            headers = {
                str(cell.value).strip().lower(): idx
                for idx, cell in enumerate(ws[1], 0)
                if cell.value
            }
            topic_col = headers.get("topic")
            status_col = headers.get("status")
            if topic_col is None or status_col is None:
                raise ValueError(f"Sheet '{sheet_name}' must have 'topic' and 'status' columns")
            for row in ws.iter_rows(min_row=2):
                if str(row[topic_col].value or "").strip() == topic.strip():
                    row[status_col].value = "Done"
                    break
            wb.save(self.file_path)


# ── Google Sheets backend (read + write live, no git involved) ─────────────────
class _SheetsBackend:
    _SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

    _RETRY = 4

    def __init__(self, sheet_id: str, sa_file: str):
        from google.oauth2.service_account import Credentials
        from googleapiclient.discovery import build

        creds = Credentials.from_service_account_file(sa_file, scopes=self._SCOPES)
        self.sheet_id = sheet_id
        self._sheets = build("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets()
        self._lock = threading.Lock()
        logger.info(f"Topics backend: Google Sheets ({sheet_id[:12]}...)")

    @staticmethod
    def _is_transient(e: Exception) -> bool:
        # SSL/socket/connection errors are all OSError subclasses; 5xx/429 carry
        # their code in the message. 4xx (403/404) are NOT transient — don't retry.
        if isinstance(e, (ssl.SSLError, socket.timeout, ConnectionError, TimeoutError, OSError)):
            return True
        s = str(e).lower()
        return any(k in s for k in (
            "eof occurred", "ssl", "timed out", "timeout", "connection reset",
            "connection aborted", "broken pipe", "temporarily unavailable",
            " 500", " 502", " 503", " 504", " 429", "rate limit",
        ))

    def _exec(self, request):
        """Run a Sheets API request with retry on transient SSL/network errors."""
        last = None
        for i in range(1, self._RETRY + 1):
            try:
                return request.execute()
            except Exception as e:
                last = e
                if i == self._RETRY or not self._is_transient(e):
                    raise
                wait = 2 ** (i - 1)
                logger.warning(
                    f"Sheets API transient error (try {i}/{self._RETRY}) — retry {wait}s: {str(e)[:120]}"
                )
                time.sleep(wait)
        raise last  # pragma: no cover

    def get_sheet_names(self) -> List[str]:
        meta = self._exec(self._sheets.get(spreadsheetId=self.sheet_id))
        return [s["properties"]["title"] for s in meta.get("sheets", [])]

    def _read_tab(self, sheet_name: str):
        resp = self._exec(self._sheets.values().get(
            spreadsheetId=self.sheet_id,
            range=f"'{sheet_name}'",
        ))
        values = resp.get("values", [])
        if not values:
            return [], []
        headers = [str(h).strip().lower() for h in values[0]]
        return headers, values[1:]

    @staticmethod
    def _cell(row, idx) -> str:
        return str(row[idx]).strip() if 0 <= idx < len(row) else ""

    def get_pending_topics(self, sheet_name: str = "Sheet1") -> List[TopicRow]:
        headers, rows = self._read_tab(sheet_name)
        if "status" not in headers:
            raise ValueError(f"Sheet '{sheet_name}' must have a 'status' column")

        def idx(name):
            return headers.index(name) if name in headers else -1

        c_status, c_topic = idx("status"), idx("topic")
        c_num, c_series, c_ep, c_pri = idx("#"), idx("series_name"), idx("episode_no"), idx("priority")

        out = []
        for row in rows:
            if self._cell(row, c_status).lower() != "pending":
                continue
            topic = self._cell(row, c_topic)
            if not topic:
                continue
            out.append(
                TopicRow(
                    number=_safe_int(self._cell(row, c_num)),
                    topic=topic,
                    series_name=self._cell(row, c_series),
                    episode_no=_safe_int(self._cell(row, c_ep)),
                    priority=self._cell(row, c_pri),
                    status=self._cell(row, c_status),
                )
            )
        out.sort(key=lambda r: r.number)
        return out

    def get_next_pending(self, sheet_name: str = "Sheet1") -> Optional[TopicRow]:
        rows = self.get_pending_topics(sheet_name=sheet_name)
        return rows[0] if rows else None

    def pending_count(self, sheet_name: str = "Sheet1") -> int:
        return len(self.get_pending_topics(sheet_name=sheet_name))

    def mark_done(self, topic: str, sheet_name: str = "Sheet1"):
        with self._lock:
            headers, rows = self._read_tab(sheet_name)
            c_topic = headers.index("topic") if "topic" in headers else -1
            c_status = headers.index("status") if "status" in headers else -1
            if c_topic < 0 or c_status < 0:
                raise ValueError(f"Sheet '{sheet_name}' must have 'topic' and 'status' columns")
            for i, row in enumerate(rows):
                if self._cell(row, c_topic) == topic.strip():
                    a1 = f"'{sheet_name}'!{_col_letter(c_status)}{i + 2}"  # +2: header + 1-based
                    self._exec(self._sheets.values().update(
                        spreadsheetId=self.sheet_id,
                        range=a1,
                        valueInputOption="RAW",
                        body={"values": [["Done"]]},
                    ))
                    logger.info(f"[{sheet_name}] marked done in Google Sheet: {topic[:50]}")
                    return
            logger.warning(f"[{sheet_name}] topic not found to mark done: {topic[:50]}")


# ── Public manager: picks backend, exposes the same interface ──────────────────
class TopicsManager:
    """Reads topics + marks them done. Uses Google Sheets when sheet_id + a valid
    service-account file are given (no git needed); otherwise the local Excel."""

    def __init__(self, file_path: str, sheet_id: str = "", sa_file: str = ""):
        use_sheets = bool(sheet_id) and bool(sa_file) and Path(sa_file).exists()
        if use_sheets:
            try:
                self._impl = _SheetsBackend(sheet_id, sa_file)
                self.is_remote = True
                return
            except Exception as e:
                logger.warning(f"Google Sheets backend init failed ({e}) — falling back to Excel")
        self._impl = _ExcelBackend(file_path)
        self.is_remote = False

    def get_sheet_names(self) -> List[str]:
        return self._impl.get_sheet_names()

    def get_pending_topics(self, sheet_name: str = "Sheet1") -> List[TopicRow]:
        return self._impl.get_pending_topics(sheet_name=sheet_name)

    def get_next_pending(self, sheet_name: str = "Sheet1") -> Optional[TopicRow]:
        return self._impl.get_next_pending(sheet_name=sheet_name)

    def pending_count(self, sheet_name: str = "Sheet1") -> int:
        return self._impl.pending_count(sheet_name=sheet_name)

    def mark_done(self, topic: str, sheet_name: str = "Sheet1"):
        return self._impl.mark_done(topic, sheet_name=sheet_name)
